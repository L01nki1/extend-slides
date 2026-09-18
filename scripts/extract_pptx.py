#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 PPTX 抽取每页文字与备注，供 extend-slides 生成扩展讲解文档使用。

用法：
    python scripts/extract_pptx.py <input.pptx> <课程输出目录> [--images]

产出：
    <课程输出目录>/<文件名>-抽取结果.md        每页文字与备注；纯图片页会被显式标注
    <课程输出目录>/assets/<页号>-img<序号>.<扩展名>   仅在传入 --images 时导出

默认**不导出图片**。按 extend-slides 的输出约定，无 OCR 时的纯图片页应当直接跳过，
而不是把课件图片贴进 Markdown 顶替正文。只有确实需要保留某张图时才加 --images。

依赖：
    pip install python-pptx
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:  # pragma: no cover - 依赖缺失时的友好提示
    sys.exit("缺少依赖 python-pptx，请先运行：pip install python-pptx")


def iter_shapes(shapes):
    """深度优先遍历形状，展开组合（GROUP）内部的形状。"""
    for shape in shapes:
        try:
            shape_type = shape.shape_type
        except Exception:  # noqa: BLE001 - 个别形状类型无法判定
            shape_type = None
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                yield from iter_shapes(shape.shapes)
            except Exception:  # noqa: BLE001
                continue
        else:
            yield shape


def shape_text(shape) -> list[str]:
    """抽取形状中的文本，含表格内容。"""
    lines: list[str] = []

    try:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    lines.append(text)
    except Exception:  # noqa: BLE001
        pass

    try:
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append("| " + " | ".join(cells) + " |")
    except Exception:  # noqa: BLE001
        pass

    return lines


def save_picture(shape, slide_no: int, index: int, assets: Path) -> str | None:
    """把图片写入 assets 目录，返回 Markdown 用的相对路径；失败返回 None。"""
    try:
        image = shape.image
    except Exception:  # noqa: BLE001
        return None

    ext = (image.ext or "png").lower()
    name = f"{slide_no:02d}-img{index:02d}.{ext}"
    try:
        (assets / name).write_bytes(image.blob)
    except OSError:
        return None
    return f"{assets.name}/{name}"


def main(argv: list[str]) -> int:
    export_images = "--images" in argv
    args = [arg for arg in argv if arg != "--images"]
    if len(args) != 2:
        print(__doc__)
        return 2

    src = Path(args[0]).expanduser()
    out_dir = Path(args[1]).expanduser()

    if not src.is_file():
        print(f"找不到文件：{src}")
        return 1
    if src.suffix.lower() not in {".pptx", ".pptm"}:
        print(f"仅支持 .pptx / .pptm，收到 {src.suffix}。若是 .ppt，请先另存为 .pptx。")
        return 1

    out_dir.mkdir(parents=True, exist_ok=True)
    assets = out_dir / "assets"
    if export_images:
        assets.mkdir(parents=True, exist_ok=True)

    try:
        prs = Presentation(str(src))
    except Exception as exc:  # noqa: BLE001
        print(f"无法打开该 PPTX：{exc}")
        return 1

    md: list[str] = [
        f"# {src.name} 抽取结果",
        "",
        "> 本文件由 extract_pptx.py 生成，仅作工作底稿，不是最终交付物。",
        "",
    ]
    slide_count = 0
    image_count = 0
    failed_images: list[str] = []
    pure_image_pages: list[int] = []
    empty_pages: list[int] = []

    for slide_no, slide in enumerate(prs.slides, start=1):
        slide_count = slide_no
        md.append(f"## 第 {slide_no} 页")
        md.append("")

        picture_count = 0
        text_lines: list[str] = []

        for shape in iter_shapes(slide.shapes):
            try:
                shape_type = shape.shape_type
            except Exception:  # noqa: BLE001
                shape_type = None

            if shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1
                if export_images:
                    rel_path = save_picture(shape, slide_no, picture_count, assets)
                    if rel_path is None:
                        failed_images.append(f"第 {slide_no} 页 图 {picture_count}")
                        md.append(f"[图片无法提取：{src.name} 第 {slide_no} 页 图 {picture_count}]")
                        md.append("")
                    else:
                        image_count += 1
                        md.append(f"![第 {slide_no} 页 图 {picture_count}]({rel_path})")
                        md.append("")
                continue

            text_lines.extend(shape_text(shape))

        if text_lines:
            md.extend(text_lines)
            md.append("")
        elif picture_count:
            pure_image_pages.append(slide_no)
            md.append(
                f"[纯图片页：无可提取文字，含 {picture_count} 张图片。"
                "按输出约定应跳过，不要把图片贴进正文；请记入交付报告的存疑条目。]"
            )
            md.append("")
        else:
            empty_pages.append(slide_no)
            md.append("[空页：既无可提取文字，也没有图片。]")
            md.append("")

        try:
            has_notes = slide.has_notes_slide
        except Exception:  # noqa: BLE001
            has_notes = False
        if has_notes:
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:  # noqa: BLE001
                notes = ""
            if notes:
                md.append(f"备注：{notes}")
                md.append("")

    out_md = out_dir / f"{src.stem}-抽取结果.md"
    out_md.write_text("\n".join(md), encoding="utf-8")

    print(f"幻灯片页数：{slide_count}")
    print(f"文字抽取结果：{out_md}")
    if export_images:
        print(f"导出图片：{image_count} 张 → {assets}")
    else:
        print("未导出图片（默认行为）。确需保留某张图时加 --images 重新运行。")
    if pure_image_pages:
        preview = "、".join(str(page) for page in pure_image_pages[:20])
        suffix = "…" if len(pure_image_pages) > 20 else ""
        print(f"纯图片页（{len(pure_image_pages)} 页，应跳过并记入交付报告）：第 {preview}{suffix} 页")
    if empty_pages:
        preview = "、".join(str(page) for page in empty_pages[:20])
        print(f"空页（{len(empty_pages)} 页）：第 {preview} 页")
    if failed_images:
        print(f"未能提取的图片（{len(failed_images)} 处）：{'、'.join(failed_images)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
